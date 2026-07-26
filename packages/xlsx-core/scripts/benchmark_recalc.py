#!/usr/bin/env python3
"""Benchmark all three recalculation engines against the SOLAR workbook.

Compares every formula cell's engine-computed value against the workbook's
original cached values (openpyxl data_only=True).  Writes the report to
dag/recalc_benchmark.json.

Each engine runs in its own subprocess with a 10-minute timeout so one hang
doesn't kill the whole benchmark.
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from xlsx_core.recalc.base import compare_values
from xlsx_core.settings import Settings


def _read_cached_values(workbook: Path) -> dict[str, object]:
    """Read all cached values (data_only) keyed by sheet!cell."""
    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
    values: dict[str, object] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    key = f"{sheet_name}!{cell.coordinate}"
                    values[key] = cell.value
    wb.close()
    return values


def _read_recalced_libreoffice(recalced_path: Path) -> dict[str, object]:
    """Read a recalculated .xlsx via openpyxl data_only."""
    wb = openpyxl.load_workbook(recalced_path, data_only=True, read_only=True)
    values: dict[str, object] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    key = f"{sheet_name}!{cell.coordinate}"
                    values[key] = cell.value
    wb.close()
    return values


def _read_recalced_formulas_pkg(values_path: Path) -> dict[str, object]:
    """Read formulas-pkg .txt output back into a dict."""
    values: dict[str, object] = {}
    if not values_path.exists():
        return values
    for line in values_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or "\t" not in line:
            continue
        key, raw_val = line.split("\t", 1)
        # Try to parse numeric values
        try:
            val: object = (
                float(raw_val)
                if "." in raw_val or "e" in raw_val.lower()
                else int(raw_val)
            )
        except ValueError:
            val = raw_val
        values[key] = val
    return values


def _read_formula_cells(workbook: Path) -> set[str]:
    """Return the set of all formula-cell canonical ids (sheet!A1)."""
    wb = openpyxl.load_workbook(workbook, data_only=False, read_only=True)
    formula_cells: set[str] = set()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.add(f"{sheet_name}!{cell.coordinate}")
    wb.close()
    return formula_cells


def _compare_engine(
    engine_result: dict,
    workbook: Path,
    cached: dict[str, object],
    formula_cells: set[str],
    temp_dir: Path,
) -> dict:
    """Compare engine output against cached values on formula cells only."""
    if not engine_result["ok"]:
        engine_result["mismatch_count"] = 0
        engine_result["total_compared"] = 0
        return engine_result

    recalced_path = engine_result.get("recalced_path")
    if not recalced_path or not Path(recalced_path).exists():
        engine_result["ok"] = False
        existing = list(engine_result.get("errors", []))
        engine_result["errors"] = [*existing, "recalced file not found"]
        engine_result["mismatch_count"] = 0
        engine_result["total_compared"] = 0
        return engine_result

    # Read recalced values
    engine_name = engine_result["engine"]
    if engine_name == "formulas-pkg":
        recalced = _read_recalced_formulas_pkg(Path(recalced_path))
    else:
        recalced = _read_recalced_libreoffice(Path(recalced_path))

    settings = Settings()
    mismatches = 0
    total = 0

    for cell_id in sorted(formula_cells):
        total += 1
        cached_val = cached.get(cell_id)
        recalced_val = recalced.get(cell_id)

        if not compare_values(
            recalced_val,
            cached_val,
            rel_tol=settings.rel_tol,
            abs_tol=settings.abs_tol,
        ):
            mismatches += 1

    engine_result["mismatch_count"] = mismatches
    engine_result["total_compared"] = total
    return engine_result


def _run_engine_in_subprocess(
    engine: str,
    workbook: Path,
    out_dir: Path,
    timeout: int = 600,
) -> dict:
    """Run a single engine benchmark in a subprocess with a timeout."""
    script = rf"""
import json, sys, tempfile
from pathlib import Path
sys.path.insert(0, r"{Path(__file__).resolve().parent.parent / "src"!s}")
from xlsx_core.recalc.base import compare_values
from xlsx_core.settings import Settings

workbook = Path(r"{workbook!s}")
out_dir = Path(r"{out_dir!s}")

if "{engine}" == "libreoffice":
    from xlsx_core.recalc.libreoffice import LibreOfficeRecalc
    adapter = LibreOfficeRecalc(Settings().libreoffice_path)
elif "{engine}" == "formulas-pkg":
    from xlsx_core.recalc.formulas_pkg import FormulasPkgRecalc
    adapter = FormulasPkgRecalc()
elif "{engine}" == "excel-com":
    from xlsx_core.recalc.excel_com import ExcelComRecalc
    adapter = ExcelComRecalc()
else:
    json.dump(
        {{"engine": "{engine}", "ok": False, "wall_seconds": 0, "errors": ["unknown"]}},
        sys.stdout,
    )
    sys.exit(0)

result = adapter.recalc(workbook, out_dir)
d = {{
    "engine": result.engine,
    "ok": result.ok,
    "wall_seconds": result.wall_seconds,
    "mismatch_count": result.mismatch_count,
    "total_compared": result.total_compared,
    "errors": list(result.errors),
    "recalced_path": str(result.recalced_path) if result.recalced_path else None,
}}
json.dump(d, sys.stdout)
"""
    script_path = out_dir / f"_bench_{engine}.py"
    script_path.write_text(script, encoding="utf-8")

    try:
        proc = subprocess.run(
            [sys.executable, str(script_path)],
            capture_output=True,
            text=True,
            timeout=timeout,
            cwd=str(Path(__file__).resolve().parent.parent),
        )
        if proc.returncode != 0:
            return {
                "engine": engine,
                "ok": False,
                "wall_seconds": 0.0,
                "mismatch_count": 0,
                "total_compared": 0,
                "errors": [f"subprocess exit {proc.returncode}: {proc.stderr.strip()}"],
                "recalced_path": None,
            }
        return json.loads(proc.stdout)
    except subprocess.TimeoutExpired:
        return {
            "engine": engine,
            "ok": False,
            "wall_seconds": float(timeout),
            "mismatch_count": 0,
            "total_compared": 0,
            "errors": [f"timeout after {timeout}s"],
            "recalced_path": None,
        }


def main() -> None:
    settings = Settings()
    workbook = settings.workbook_path.resolve()
    if not workbook.exists():
        print(f"Workbook not found: {workbook}", file=sys.stderr)
        sys.exit(1)

    # Output dirs
    dag_dir = Path("dag")
    dag_dir.mkdir(parents=True, exist_ok=True)

    bench_tmp = Path(tempfile.mkdtemp(prefix="recalc_bench_"))

    # Read cached values and formula cell list once
    print("Reading cached values from workbook...")
    cached = _read_cached_values(workbook)
    formula_cells = _read_formula_cells(workbook)
    print(f"  {len(cached)} non-blank cells, {len(formula_cells)} formula cells")

    engines = ["libreoffice", "formulas-pkg", "excel-com"]
    results: list[dict] = []

    for engine in engines:
        print(f"\n--- {engine} ---")
        engine_out = bench_tmp / engine
        engine_out.mkdir(parents=True, exist_ok=True)

        result = _run_engine_in_subprocess(engine, workbook, engine_out, timeout=600)
        result = _compare_engine(result, workbook, cached, formula_cells, engine_out)
        results.append(result)

        status = "OK" if result["ok"] else "FAILED"
        mm = result.get("mismatch_count", 0)
        tot = result.get("total_compared", 0)
        print(f"  {status}  wall={result['wall_seconds']:.1f}s  mismatches={mm}/{tot}")

    # Clean up temp
    import shutil

    shutil.rmtree(bench_tmp, ignore_errors=True)

    # Write report
    report_path = dag_dir / "recalc_benchmark.json"
    report = {
        "engines": results,
        "workbook": str(workbook),
        "formula_cells": len(formula_cells),
    }
    report_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"\nReport written to {report_path}")

    # Print summary
    print("\n=== BENCHMARK SUMMARY ===")
    for r in results:
        print(
            f"  {r['engine']:15s}  ok={r['ok']!s:5s}  "
            f"wall={r['wall_seconds']:8.2f}s  "
            f"mismatches={r.get('mismatch_count', 0)}/"
            f"{r.get('total_compared', 0)}"
            f"  errors={r['errors']}",
        )


if __name__ == "__main__":
    main()
