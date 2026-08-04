"""Tests for domain schema/compute/download endpoints."""
from __future__ import annotations

from fastapi.testclient import TestClient
from solar_server.main import app

client = TestClient(app)


# ---------------------------------------------------------------------------
# GET /api/schema
# ---------------------------------------------------------------------------


def test_api_schema_returns_domains() -> None:
    """GET /api/schema returns domain list with items."""
    r = client.get("/api/schema")
    assert r.status_code == 200
    data = r.json()
    assert "domains" in data
    domains = data["domains"]
    # params, invest, debt, cost, pnl, cashflow, finplan, balance, valuation
    assert len(domains) == 9
    keys = {d["key"] for d in domains}
    assert keys >= {"params", "pnl", "valuation", "cashflow"}

    # Spot-check pnl schema
    pnl = next(d for d in domains if d["key"] == "pnl")
    assert pnl["sheet"] == "损益"
    assert pnl["label"] == "损益表 (PnL)"
    assert len(pnl["items"]) > 5
    sales = next(it for it in pnl["items"] if it["key"] == "sales_revenue")
    assert sales["label"] == "销售收入"
    assert "power_generation" in sales["formula"]


# ---------------------------------------------------------------------------
# POST /api/domains/compute
# ---------------------------------------------------------------------------


def test_domains_compute_v2_returns_cells() -> None:
    """POST /api/domains/compute v2 returns cells with values."""
    r = client.post("/api/domains/compute", json={"version": "v2"})
    assert r.status_code == 200
    data = r.json()
    assert data["version"] == "v2"
    assert len(data["domains"]) == 9

    # Spot-check pnl domain
    pnl = next(d for d in data["domains"] if d["key"] == "pnl")
    assert pnl["sheet"] == "损益"
    assert len(pnl["items"]) > 5

    # Each item should have cells
    for item in pnl["items"]:
        assert "cells" in item
        assert "key" in item
        assert "label" in item
        assert "unit" in item
        assert "formula" in item

    # The sales_revenue item should have numeric cells
    sales = next(it for it in pnl["items"] if it["key"] == "sales_revenue")
    assert len(sales["cells"]) >= 1
    # First operating year (2021) sales_revenue = ~5645.35 (万元)
    first_val = sales["cells"][0]["value"]
    assert isinstance(first_val, (int, float))
    assert abs(first_val - 5645.35030420354) < 1e-5


def test_domains_compute_v2_response_shape() -> None:
    """POST /api/domains/compute v2 preserves app.js-expected shape.

    The frontend expects: domains[].{key, sheet, label, depends_on, items[]}
    items[].{key, label, unit, formula, inputs, cells[]}
    cells[].{col, row, value}
    """
    r = client.post("/api/domains/compute", json={"version": "v2"})
    assert r.status_code == 200
    data = r.json()

    for domain in data["domains"]:
        # Required domain fields
        assert "key" in domain
        assert "sheet" in domain
        assert "label" in domain
        assert "depends_on" in domain
        assert "items" in domain
        assert isinstance(domain["items"], list)

        for item in domain["items"]:
            # Required item fields (matched by app.js renderDomainTables / renderPanel)
            assert "key" in item
            assert "label" in item
            assert "unit" in item
            assert "formula" in item
            assert "cells" in item
            assert isinstance(item["cells"], list)

            for cell in item["cells"]:
                assert "col" in cell
                assert "row" in cell
                assert "value" in cell
                # col is a string (year or column letter)
                assert isinstance(cell["col"], str)
                # row is an int
                assert isinstance(cell["row"], int)


def test_domains_compute_v1_works() -> None:
    """POST /api/domains/compute v1 also works."""
    r = client.post("/api/domains/compute", json={"version": "v1"})
    assert r.status_code == 200
    data = r.json()
    assert len(data["domains"]) >= 1


def test_domains_compute_v3_422() -> None:
    """POST /api/domains/compute v3 returns 422."""
    r = client.post("/api/domains/compute", json={"version": "v3"})
    assert r.status_code == 422
    assert "v3" in r.json()["detail"]


def test_domains_compute_defaults_to_v2() -> None:
    """POST /api/domains/compute without body defaults to v2."""
    r = client.post("/api/domains/compute")
    assert r.status_code == 200
    assert r.json()["version"] == "v2"


def test_domains_compute_v2_inputs_override() -> None:
    """POST /api/domains/compute v2 with inputs override changes output."""
    r_default = client.post("/api/domains/compute", json={"version": "v2"})
    assert r_default.status_code == 200

    r_override = client.post(
        "/api/domains/compute",
        json={"version": "v2", "inputs": {"installed_capacity_mw": 30.0}},
    )
    assert r_override.status_code == 200

    # With 30MW vs 15MW default, power_generation doubles
    default_pnl = next(d for d in r_default.json()["domains"] if d["key"] == "pnl")
    override_pnl = next(d for d in r_override.json()["domains"] if d["key"] == "pnl")

    default_sales = next(
        it for it in default_pnl["items"] if it["key"] == "power_generation"
    )
    override_sales = next(
        it for it in override_pnl["items"] if it["key"] == "power_generation"
    )

    ratio = override_sales["cells"][0]["value"] / default_sales["cells"][0]["value"]
    assert abs(ratio - 2.0) < 0.01


# ---------------------------------------------------------------------------
# GET /api/download/domains.xlsx
# ---------------------------------------------------------------------------


def test_download_xlsx_returns_valid_file() -> None:
    """GET /api/download/domains.xlsx returns valid xlsx with expected sheets."""
    r = client.get("/api/download/domains.xlsx", params={"version": "v2"})
    assert r.status_code == 200
    ct = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r.headers["content-type"] == ct

    # Read back with openpyxl
    import io

    import openpyxl

    buf = io.BytesIO(r.content)
    wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    # Should have domain sheets (中文 names)
    assert len(sheet_names) >= 8
    assert "损益" in sheet_names
    assert "成本" in sheet_names
    assert "估值结果" in sheet_names

    # Spot-check: 损益 sheet header
    ws = wb["损益"]
    rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
    header = [str(c) if c else "" for c in rows[0]]
    assert header[:3] == ["项目", "公式", "单位"]
    wb.close()


def test_download_xlsx_v1_works() -> None:
    """GET /api/download/domains.xlsx v1 returns 200."""
    r = client.get("/api/download/domains.xlsx", params={"version": "v1"})
    assert r.status_code == 200


def test_download_xlsx_v3_422() -> None:
    """GET /api/download/domains.xlsx v3 returns 422."""
    r = client.get("/api/download/domains.xlsx", params={"version": "v3"})
    assert r.status_code == 422
    assert "v3" in r.json()["detail"]


def test_download_xlsx_freezes_header() -> None:
    """XLSX sheet has frozen top row."""
    import io

    import openpyxl

    r = client.get("/api/download/domains.xlsx", params={"version": "v2"})
    assert r.status_code == 200
    buf = io.BytesIO(r.content)
    wb = openpyxl.load_workbook(buf)
    ws = wb["损益"]
    assert ws.freeze_panes == "A2"
    wb.close()


# ---------------------------------------------------------------------------
# Static mount — gracefully handled if dir missing
# ---------------------------------------------------------------------------


def test_root_serves_or_falls_through() -> None:
    """GET / returns something (static or fallback — not a crash)."""
    r = client.get("/")
    # If static dir is present, returns 200; else 404
    assert r.status_code in (200, 404)
