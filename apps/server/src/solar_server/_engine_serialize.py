"""Serialize ModelResults for HTTP responses.

Provides domain-structured JSON for /compute/v2, DomainsComputeResponse
for /api/domains/compute (preserving the app.js-expected shape), and an
openpyxl workbook builder for /api/download/domains.xlsx.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
import polars as pl
from solar_v2.results import ModelResults
from solar_v2.schema import all_domains

from solar_server.schemas import (
    CellValue,
    DomainComputeResponse,
    DomainsComputeResponse,
    ItemCell,
    scalar_to_json,
)

# ---------------------------------------------------------------------------
# Column-letter helpers (matching domains.py)
# ---------------------------------------------------------------------------


def _col_index_to_letter(idx: int) -> str:
    """0 -> 'A', 1 -> 'B', ..."""
    result = ""
    n = idx
    while n >= 0:
        result = chr(ord("A") + n % 26) + result
        n = n // 26 - 1
    return result


# ---------------------------------------------------------------------------
# Domain JSON for /compute/v2
# ---------------------------------------------------------------------------


def model_results_to_domain_json(results: ModelResults) -> dict[str, object]:
    """Serialize ModelResults to domain-structured JSON.

    Returns::

        {
          "version": "v2",
          "params": { "key": "params", "sheet": "参数表", ...,
            "items": [{"key": ..., "label": ..., "unit": ..., "value": ...}, ...]
          },
          ...
          "headline": { "equity_irr": ..., ... }
        }
    """
    schemas = all_domains()

    domains: dict[str, object] = {}
    for ds in schemas:
        domain_result = getattr(results, ds.key)
        items: list[dict[str, object]] = []
        for item_schema in ds.items:
            entry: dict[str, object] = {
                "key": item_schema.key,
                "label": item_schema.label,
                "unit": item_schema.unit,
                "formula": item_schema.formula,
                "kind": item_schema.kind,
            }
            if item_schema.kind == "scalar":
                entry["value"] = _read_scalar(domain_result, item_schema.key)
            else:
                series = _read_series(domain_result, item_schema.key)
                if series is not None:
                    entry["values"] = _series_to_list(series)
                else:
                    entry["values"] = None
            items.append(entry)

        years = _domain_years(domain_result)
        domains[ds.key] = {
            "key": ds.key,
            "sheet": ds.sheet,
            "label": ds.label,
            "years": years,
            "items": items,
        }

    return {
        "version": "v2",
        **domains,
        "headline": {
            "equity_irr": results.equity_irr,
            "project_irr_after_tax": results.project_irr_after_tax,
            "equity_npv": results.equity_npv,
            "project_npv_after_tax": results.project_npv_after_tax,
            "equity_sale_price": results.equity_sale_price,
        },
    }


# ---------------------------------------------------------------------------
# DomainsComputeResponse for /api/domains/compute  (app.js shape)
# ---------------------------------------------------------------------------


def model_results_to_compute_response(
    results: ModelResults,
) -> DomainsComputeResponse:
    """Build DomainsComputeResponse from ModelResults + domain schemas.

    Preserves the exact response shape consumed by ``apps/server/static/app.js``:
    ``domains[].items[].cells[]`` with ``col``, ``row``, ``value`` —
    each series item gets one cell per operating year, column letters C–AA
    mapping to calendar years; each scalar item gets a single cell in column C.

    A synthetic ``year_labels`` item is included per domain (when the domain
    has series items) so the frontend detail panel can render year headers.
    """
    schemas = all_domains()
    domains: list[DomainComputeResponse] = []

    for ds in schemas:
        domain_result = getattr(results, ds.key)
        years = _domain_years(domain_result)  # operating-year calendar years

        items: list[ItemCell] = []

        # If there are series items, emit a year_labels row
        has_series = any(it.kind == "series" for it in ds.items)
        if has_series and years:
            year_cells: list[CellValue] = []
            for i, y in enumerate(years):
                year_cells.append(
                    CellValue(
                        col=_col_index_to_letter(i + 2),  # start at column C
                        row=0,
                        value=y,
                    )
                )
            items.append(
                ItemCell(
                    key="year_labels",
                    label="年份",
                    unit="-",
                    formula="运营年度",
                    inputs=(),
                    cells=tuple(year_cells),
                )
            )

        for item_schema in ds.items:
            cells: list[CellValue] = []
            row_base = item_schema.rows[0] if item_schema.rows else 1

            if item_schema.kind == "scalar":
                val = _read_scalar(domain_result, item_schema.key)
                cells.append(
                    CellValue(
                        col="C",
                        row=row_base,
                        value=scalar_to_json(val),
                    )
                )
            else:
                series = _read_series(domain_result, item_schema.key)
                if series is not None and years:
                    for i, val in enumerate(_series_to_list(series)):
                        cells.append(
                            CellValue(
                                col=_col_index_to_letter(i + 2),
                                row=row_base,
                                value=scalar_to_json(val),
                            )
                        )

            items.append(
                ItemCell(
                    key=item_schema.key,
                    label=item_schema.label,
                    unit=item_schema.unit,
                    formula=item_schema.formula,
                    inputs=item_schema.inputs,
                    cells=tuple(cells),
                )
            )

        domains.append(
            DomainComputeResponse(
                key=ds.key,
                sheet=ds.sheet,
                label=ds.label,
                depends_on=ds.depends_on,
                items=tuple(items),
            )
        )

    return DomainsComputeResponse(version="v2", domains=tuple(domains))


# ---------------------------------------------------------------------------
# XLSX builder for /api/download/domains.xlsx
# ---------------------------------------------------------------------------


def model_results_to_xlsx(results: ModelResults) -> io.BytesIO:
    """Build an openpyxl workbook with one sheet per domain from ModelResults.

    Sheet layout:
    - Row 1 (header): 项目 | 公式 | 单位 | year columns...
    - Row 2 (年份): 年份 | | | year values...
    - Subsequent rows: one per domain item, label | formula | unit | values...
    """
    schemas = all_domains()
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    if schemas:
        wb.remove(default_sheet)

    for ds in schemas:
        domain_result = getattr(results, ds.key)
        years = _domain_years(domain_result)
        n_years = len(years) if years else 0

        # Column letters C onwards
        col_letters = [_col_index_to_letter(i + 2) for i in range(n_years)]

        # Sheet name from schema
        ws_title = ds.sheet[:31]
        ws = wb.create_sheet(title=ws_title)

        # Header row
        header = ["项目", "公式", "单位", *col_letters]
        ws.append(header)

        # Year row
        if years:
            year_row: list[object] = ["年份", "", ""]
            year_row.extend(years)
            ws.append(year_row)

        # Data rows
        for item_schema in ds.items:
            row_data: list[object] = [
                item_schema.label,
                item_schema.formula,
                item_schema.unit,
            ]

            if item_schema.kind == "scalar":
                val = _read_scalar(domain_result, item_schema.key)
                row_data.append(_xlsx_value(val))
                # Pad remaining columns
                row_data.extend([""] * (n_years - 1))
            else:
                series = _read_series(domain_result, item_schema.key)
                if series is not None:
                    vals = _series_to_list(series)
                    for v in vals:
                        row_data.append(_xlsx_value(v))
                else:
                    row_data.extend([""] * n_years)

            ws.append(row_data)

        # Column widths
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 8
        for i in range(n_years):
            col_letter = _col_index_to_letter(i + 3)  # D onwards
            ws.column_dimensions[col_letter].width = 12

        # Header styling
        from openpyxl.styles import Font as OpFont

        bold_font = OpFont(bold=True)
        for cell in ws[1]:
            cell.font = bold_font
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Internal helpers — read from typed domain results
# ---------------------------------------------------------------------------


def _read_scalar(domain_result: object, key: str) -> Any:
    """Read a scalar value from a domain result.

    Tries: exact key match, suffix variants, prefix match with known
    suffixes, then a fuzzy search over all scalar fields.
    """
    from dataclasses import fields, is_dataclass

    scalars = getattr(domain_result, "scalars", None)
    if scalars is not None:
        if hasattr(scalars, key):
            return getattr(scalars, key)
        # Try common suffixes
        for suffix in ("_after_tax", "_pre_tax", "_total", "_pct"):
            alt = key + suffix
            if hasattr(scalars, alt):
                return getattr(scalars, alt)
        # Try stripping _period, _years etc and adding suffixes
        for strip in ("_period", "_years", "_ratio", "_rate", "_pct", "_total"):
            if key.endswith(strip):
                base = key[: -len(strip)]
                for suffix in ("_after_tax", "_pre_tax", "", "_total"):
                    alt = base + suffix
                    if hasattr(scalars, alt):
                        return getattr(scalars, alt)
        # Fuzzy: find field whose name starts with the key
        if is_dataclass(scalars):
            for field in fields(scalars):  # type: ignore[arg-type]
                if field.name.startswith(key):
                    return getattr(scalars, field.name)

    if hasattr(domain_result, key):
        return getattr(domain_result, key)

    raise KeyError(f"Scalar {key!r} not found on {type(domain_result).__name__}")

def _read_series(domain_result: object, key: str) -> pl.Series | None:
    """Read a year-series column from a domain result's frame(s)."""
    for attr in ("frame", "income", "market"):
        frame = getattr(domain_result, attr, None)
        if frame is not None and key in frame.columns:
            return frame[key]
    return None


def _series_to_list(series: pl.Series) -> list[float | None]:
    """Convert polars Series to plain Python list, replacing inf/nan with None."""
    result: list[float | None] = []
    for v in series.to_list():
        if isinstance(v, float) and (v != v or v == float("inf") or v == float("-inf")):
            result.append(None)
        else:
            result.append(v)
    return result


def _domain_years(domain_result: object) -> list[int]:
    """Extract the calendar years from a domain result.

    Returns the ``years`` attribute if present, otherwise extracts from
    the ``frame["year"]`` column, falling back to empty list.
    """
    # Most domain results have a `years` attribute
    years = getattr(domain_result, "years", None)
    if years is not None:
        return list(years)

    # invest uses `periods` not `years`
    periods = getattr(domain_result, "periods", None)
    if periods is not None:
        return list(periods)

    # valuation has .income, .market frames with year column
    for attr in ("income", "frame"):
        frame = getattr(domain_result, attr, None)
        if frame is not None and "year" in frame.columns:
            return [int(y) for y in frame["year"].to_list()]

    return []


def _xlsx_value(val: object) -> object:
    """Convert a value to something openpyxl can write."""
    import math

    if val is None:
        return ""
    if isinstance(val, float) and (math.isnan(val) or math.isinf(val)):
        return ""
    if isinstance(val, (int, float, str)):
        return val
    return str(val)
