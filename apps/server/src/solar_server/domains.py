"""Domain schema and compute endpoints for solar-server.

GET /api/schema        — return DomainSchema list from solar-v2
POST /api/domains/compute — compute values per domain item
GET /api/download/domains.xlsx — Excel export with one sheet per domain
"""

from __future__ import annotations

import io
import logging
from typing import cast

import openpyxl
from fastapi import Request
from solar_v2.schema import DomainSchema
from xlsx_core.model import Scalar
from xlsx_core.settings import Settings
from xlsx_core.yamlutil import load_dag_doc

from solar_server.schemas import (
    CellValue,
    DomainComputeResponse,
    DomainSchemaResponse,
    DomainsComputeResponse,
    DomainsSchemaResponse,
    ItemCell,
    ItemSchemaResponse,
    scalar_to_json,
)

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# DAG node index — loaded once and stored on app.state
# ---------------------------------------------------------------------------


class _NodeInfo:
    """Lightweight DAG node record for cell resolution."""

    __slots__ = ("col", "node_id", "row", "sheet", "type")

    def __init__(self, node_id: str, sheet: str, col: str, row: int, type: str) -> None:
        self.node_id = node_id
        self.sheet = sheet
        self.col = col
        self.row = row
        self.type = type


def _load_dag_nodes(dag_path: str) -> list[_NodeInfo]:
    """Parse dag YAML and return formula/error nodes."""
    with open(dag_path, encoding="utf-8") as fh:
        dag = load_dag_doc(fh.name)
    nodes: list[_NodeInfo] = []
    for n in dag["nodes"]:
        if n["type"] in ("formula", "error"):
            nodes.append(
                _NodeInfo(
                    node_id=n["id"],
                    sheet=n["sheet"],
                    col=n["col"],
                    row=n["row"],
                    type=n["type"],
                )
            )
    return nodes


def _ensure_dag_index(request: Request, settings: Settings) -> list[_NodeInfo]:
    """Get or build the DAG node index, cached on app.state."""
    dag_path = str(settings.dag_path)
    cache = getattr(request.app.state, "_dag_index_cache", None)
    if cache is not None and cache.get("path") == dag_path:
        return cast(list[_NodeInfo], cache["nodes"])
    nodes = _load_dag_nodes(dag_path)
    request.app.state._dag_index_cache = {"path": dag_path, "nodes": nodes}
    return nodes


# ---------------------------------------------------------------------------
# Schema loading — per-module defensive import
# ---------------------------------------------------------------------------

_DOMAIN_MODULES: tuple[str, ...] = (
    "solar_v2.domains.params",
    "solar_v2.domains.invest",
    "solar_v2.domains.debt",
    "solar_v2.domains.cost",
    "solar_v2.domains.pnl",
    "solar_v2.domains.cashflow",
    "solar_v2.domains.finplan",
    "solar_v2.domains.balance",
    "solar_v2.domains.valuation",
)


def _load_all_domains() -> list[DomainSchema]:
    """Import each domain module and collect SCHEMA where available.

    Modules without SCHEMA are skipped without blocking others, so partial
    schema coverage is tolerated.
    """
    raw: list[DomainSchema] = []
    for mod_name in _DOMAIN_MODULES:
        try:
            mod = __import__(mod_name, fromlist=["SCHEMA"])
            schema = getattr(mod, "SCHEMA", None)
            if isinstance(schema, DomainSchema):
                raw.append(schema)
        except Exception:
            logger.debug(
                "Failed to load schema from %s", mod_name, exc_info=True
            )
    return raw


def _convert_domain_schema(ds: DomainSchema) -> DomainSchemaResponse:
    """Convert a solar_v2 DomainSchema to the response model."""
    items: list[ItemSchemaResponse] = []
    for it in ds.items:
        items.append(
            ItemSchemaResponse(
                key=it.key,
                label=it.label,
                unit=it.unit,
                formula=it.formula,
                inputs=it.inputs,
                rows=it.rows,
            )
        )
    return DomainSchemaResponse(
        key=ds.key,
        sheet=ds.sheet,
        label=ds.label,
        depends_on=ds.depends_on,
        items=tuple(items),
    )


# ---------------------------------------------------------------------------
# Schema endpoint
# ---------------------------------------------------------------------------


def _build_schema_response() -> DomainsSchemaResponse:
    """Build the schema response from solar-v2 domain schemas."""
    raw = _load_all_domains()
    domains = tuple(_convert_domain_schema(ds) for ds in raw)
    return DomainsSchemaResponse(domains=domains)


# ---------------------------------------------------------------------------
# Domains compute endpoint
# ---------------------------------------------------------------------------


def _build_item_cells(
    item: ItemSchemaResponse,
    sheet: str,
    dag_nodes: list[_NodeInfo],
    cached_values: dict[str, object],
) -> ItemCell:
    """Build an ItemCell by resolving item rows against DAG nodes and values."""
    row_set = set(item.rows)
    cells: list[CellValue] = []
    for dn in dag_nodes:
        if dn.sheet == sheet and dn.row in row_set:
            val = cached_values.get(dn.node_id)
            cells.append(
                CellValue(
                    node=dn.node_id,
                    col=dn.col,
                    row=dn.row,
                    value=scalar_to_json(cast(Scalar, val)),
                )
            )
    # sort by row then column letter
    def _col_sort_key(c: CellValue) -> tuple[int, str]:
        return (c.row, c.col)

    cells.sort(key=_col_sort_key)
    return ItemCell(
        key=item.key,
        label=item.label,
        unit=item.unit,
        formula=item.formula,
        inputs=item.inputs,
        cells=tuple(cells),
    )


def _build_domains_compute(
    version: str,
    dag_nodes: list[_NodeInfo],
    cached_values: dict[str, object],
) -> DomainsComputeResponse:
    """Build the domains compute response."""
    raw = _load_all_domains()

    domains: list[DomainComputeResponse] = []
    for ds in raw:
        items: list[ItemCell] = []
        for it in ds.items:  # type: ignore[attr-defined]
            sr = ItemSchemaResponse(
                key=it.key,
                label=it.label,
                unit=it.unit,
                formula=it.formula,
                inputs=it.inputs,
                rows=it.rows,
            )
            items.append(_build_item_cells(sr, ds.sheet, dag_nodes, cached_values))  # type: ignore[attr-defined]
        domains.append(
            DomainComputeResponse(
                key=ds.key,  # type: ignore[attr-defined]
                label=ds.label,  # type: ignore[attr-defined]
                sheet=ds.sheet,  # type: ignore[attr-defined]
                depends_on=ds.depends_on,  # type: ignore[attr-defined]
                items=tuple(items),
            )
        )
    return DomainsComputeResponse(version=version, domains=tuple(domains))


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

# Max columns reasonable for our year-series sheets
_MAX_YEAR_COLUMNS = 30


def _col_index_to_letter(idx: int) -> str:
    """0 -> 'A', 1 -> 'B', ..., 25 -> 'Z', 26 -> 'AA', ..."""
    result = ""
    n = idx
    while n >= 0:
        result = chr(ord("A") + n % 26) + result
        n = n // 26 - 1
    return result


def _col_letter_to_index(letter: str) -> int:
    """'A' -> 0, 'B' -> 1, ..., 'Z' -> 25, 'AA' -> 26, ..."""
    idx = 0
    for ch in letter:
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx - 1


def _gather_col_letters(
    sheet: str,
    dag_nodes: list[_NodeInfo],
    item_rows: set[int],
) -> list[str]:
    """Collect unique column letters used by formula/error nodes for item rows."""
    seen: set[str] = set()
    for dn in dag_nodes:
        if dn.sheet == sheet and dn.row in item_rows:
            seen.add(dn.col)
    # Sort by column index
    return sorted(seen, key=_col_letter_to_index)


def _search_year_labels(
    items: tuple[ItemSchemaResponse, ...],
    sheet: str,
    dag_nodes: list[_NodeInfo],
    cached: dict[str, object],
) -> dict[str, str] | None:
    """Find the year_labels item and return {col: year_value_str} or None."""
    for it in items:
        if it.key == "year_labels":
            row_set = set(it.rows)
            years: dict[str, str] = {}
            for dn in dag_nodes:
                if (
                    dn.sheet == sheet
                    and dn.row in row_set
                    and dn.type in ("formula", "error")
                ):
                    val = cached.get(dn.node_id)
                    if isinstance(val, (int, float)):
                        years[dn.col] = str(int(val))
                    elif isinstance(val, str):
                        years[dn.col] = val
            if years:
                return years
    return None


def _build_xlsx(
    dag_nodes: list[_NodeInfo],
    cached_values: dict[str, object],
) -> io.BytesIO:
    """Build an openpyxl workbook with one sheet per domain."""
    raw = _load_all_domains()

    wb = openpyxl.Workbook()
    # Remove default sheet; we create our own
    default_sheet = wb.active
    if raw:
        wb.remove(default_sheet)

    for ds in raw:
        items = tuple(
            ItemSchemaResponse(
                key=it.key, label=it.label, unit=it.unit,
                formula=it.formula, inputs=it.inputs, rows=it.rows,
            )
            for it in ds.items
        )

        # Collect all column letters used by formula/error rows
        all_item_rows: set[int] = set()
        for it in items:
            all_item_rows.update(it.rows)
        col_letters = _gather_col_letters(ds.sheet, dag_nodes, all_item_rows)

        # Search for year labels
        year_labels = _search_year_labels(items, ds.sheet, dag_nodes, cached_values)

        # Create sheet
        ws_title = ds.sheet[:31]  # Excel's limit
        ws = wb.create_sheet(title=ws_title)

        # Row 1: header — 项目 | 公式 | 单位 | col_letters...
        header = ["项目", "公式", "单位", *col_letters]
        ws.append(header)

        # Row 2: 年份 row if we have year labels
        if year_labels:
            year_row = ["年份", "", ""]
            for cl in col_letters:
                year_row.append(year_labels.get(cl, ""))
            ws.append(year_row)

        # Data rows: one per item
        for it in items:
            row_set = set(it.rows)
            # Build value map: col_letter → value_json
            col_vals: dict[str, object] = {}
            for dn in dag_nodes:
                if dn.sheet == ds.sheet and dn.row in row_set:
                    val = cached_values.get(dn.node_id)
                    if isinstance(val, (int, float)):
                        col_vals[dn.col] = val
                    else:
                        col_vals[dn.col] = scalar_to_json(cast(Scalar, val))

            # For the value column, display ErrorMarker error strings
            row_data = [it.label, it.formula, it.unit]
            for cl in col_letters:
                v = col_vals.get(cl)
                if v is None:
                    row_data.append("")
                elif isinstance(v, (int, float)):
                    row_data.append(v)
                else:
                    # ErrorMarker or other — use the string representation
                    if hasattr(v, "error"):
                        row_data.append(v.error)  # type: ignore[union-attr]
                    else:
                        row_data.append(str(v) if v is not None else "")
            ws.append(row_data)

        # Column widths: reasonable defaults
        ws.column_dimensions["A"].width = 22  # 项目
        ws.column_dimensions["B"].width = 45  # 公式
        ws.column_dimensions["C"].width = 8   # 单位
        for i, _cl in enumerate(col_letters, start=4):
            ws.column_dimensions[_col_index_to_letter(i - 1)].width = 12

        # Header styling — bold top row, freeze
        from openpyxl.styles import Font
        bold_font = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold_font
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
